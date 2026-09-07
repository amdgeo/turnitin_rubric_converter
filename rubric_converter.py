"""Turnitin rubric (.rbc) parsing and export functionality."""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


class RubricConversionError(ValueError):
    """Raised when a rubric cannot be parsed, validated, or exported."""


@dataclass(frozen=True)
class RubricTable:
    headers: list[str]
    rows: list[list[str]]
    criteria: list[dict[str, Any]]
    scales: list[dict[str, Any]]
    metadata: dict[str, Any]


def load_rubric(path: str | Path) -> dict[str, Any]:
    """Load and validate an .rbc file, which contains JSON data."""
    rubric_path = Path(path)
    if not rubric_path.is_file():
        raise RubricConversionError(f"Input file does not exist: {rubric_path}")
    try:
        with rubric_path.open("r", encoding="utf-8-sig") as source:
            data = json.load(source)
    except (OSError, json.JSONDecodeError) as exc:
        raise RubricConversionError(f"Could not read rubric '{rubric_path}': {exc}") from exc
    validate_rubric(data)
    return data


def validate_rubric(data: Any) -> None:
    """Validate the sections and references required by the Turnitin format."""
    if not isinstance(data, dict):
        raise RubricConversionError("The rubric root must be a JSON object.")
    required = ("RubricScale", "RubricCriterion", "RubricCriterionScale")
    for section in required:
        if not isinstance(data.get(section), list):
            raise RubricConversionError(f"Missing or invalid '{section}' section.")

    _validate_unique_ids(data["RubricScale"], "RubricScale")
    _validate_unique_ids(data["RubricCriterion"], "RubricCriterion")
    scale_ids = {item["id"] for item in data["RubricScale"]}
    criterion_ids = {item["id"] for item in data["RubricCriterion"]}
    for index, scale in enumerate(data["RubricScale"], 1):
        _require_text(scale, "name", f"RubricScale item {index}")
    for index, criterion in enumerate(data["RubricCriterion"], 1):
        if not any(criterion.get(key) is not None for key in ("description", "name")):
            raise RubricConversionError(
                f"RubricCriterion item {index} must have a description or name."
            )
    for index, entry in enumerate(data["RubricCriterionScale"], 1):
        if not isinstance(entry, dict):
            raise RubricConversionError(f"RubricCriterionScale item {index} is not an object.")
        if entry.get("criterion") not in criterion_ids:
            raise RubricConversionError(
                f"RubricCriterionScale item {index} references an unknown criterion."
            )
        if entry.get("scale_value") not in scale_ids:
            raise RubricConversionError(
                f"RubricCriterionScale item {index} references an unknown scale."
            )
        if "description" not in entry:
            raise RubricConversionError(
                f"RubricCriterionScale item {index} is missing description."
            )


def _validate_unique_ids(items: Iterable[dict[str, Any]], section: str) -> None:
    seen: set[Any] = set()
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict) or item.get("id") is None or item.get("id") == "":
            raise RubricConversionError(f"{section} item {index} is missing an id.")
        if item["id"] in seen:
            raise RubricConversionError(f"{section} contains duplicate id '{item['id']}'.")
        seen.add(item["id"])


def _require_text(item: dict[str, Any], key: str, label: str) -> None:
    if not isinstance(item.get(key), str) or not item[key].strip():
        raise RubricConversionError(f"{label} is missing a non-empty '{key}'.")


def build_table(data: dict[str, Any], use_name_and_value: bool = False) -> RubricTable:
    """Build the editable matrix used by CSV and the primary Excel sheet."""
    validate_rubric(data)
    scales = data["RubricScale"]
    criteria = data["RubricCriterion"]
    headers = ["Criteria"] + [str(scale["name"]) for scale in scales]
    rows: list[list[str]] = []
    for criterion in criteria:
        if use_name_and_value:
            name = str(criterion.get("name", criterion.get("description", "")))
            value = criterion.get("value")
            label = f"{name} ({value})" if value is not None else name
        else:
            label = str(criterion.get("description", criterion.get("name", "")))
        rows.append([label] + [""] * len(scales))

    criterion_rows = {criterion["id"]: index for index, criterion in enumerate(criteria)}
    scale_columns = {scale["id"]: index + 1 for index, scale in enumerate(scales)}
    for entry in data["RubricCriterionScale"]:
        rows[criterion_rows[entry["criterion"]]][scale_columns[entry["scale_value"]]] = str(
            entry.get("description", "")
        )
    metadata = {
        key: value
        for key, value in data.items()
        if key not in {"RubricScale", "RubricCriterion", "RubricCriterionScale"}
    }
    return RubricTable(headers, rows, criteria, scales, metadata)


def _safe_csv_value(value: Any) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def write_csv(table: RubricTable, output_path: str | Path) -> None:
    """Write the rubric matrix as UTF-8 CSV with spreadsheet formula protection."""
    try:
        with Path(output_path).open("w", newline="", encoding="utf-8") as target:
            writer = csv.writer(target)
            writer.writerow(table.headers)
            writer.writerows([[_safe_csv_value(value) for value in row] for row in table.rows])
    except OSError as exc:
        raise RubricConversionError(f"Could not write CSV '{output_path}': {exc}") from exc


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width


def _excel_safe_value(value: Any) -> Any:
    """Convert values that openpyxl cannot write directly (lists, dicts) to text."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (list, dict, tuple, set)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def write_excel(table: RubricTable, output_path: str | Path) -> None:
    """Write a formatted workbook containing matrix, details, scales, and metadata."""
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]
    _write_sheet(workbook, "Criteria Matrix", table.headers, table.rows)
    criterion_keys = sorted({key for item in table.criteria for key in item})
    _write_sheet(
        workbook,
        "Criteria Details",
        criterion_keys,
        [[_excel_safe_value(item.get(key, "")) for key in criterion_keys] for item in table.criteria],
    )
    scale_keys = sorted({key for item in table.scales for key in item})
    _write_sheet(
        workbook,
        "Scale Values",
        scale_keys,
        [[_excel_safe_value(item.get(key, "")) for key in scale_keys] for item in table.scales],
    )
    if table.metadata:
        _write_sheet(
            workbook,
            "Metadata",
            ["Field", "Value"],
            [[key, _excel_safe_value(value)] for key, value in table.metadata.items()],
        )
    try:
        workbook.save(output_path)
    except OSError as exc:
        raise RubricConversionError(f"Could not write Excel file '{output_path}': {exc}") from exc


def convert_rubric(
    input_path: str | Path,
    output_path: str | Path,
    output_format: str,
    use_name_and_value: bool = False,
) -> RubricTable:
    """Convert one rubric and return its generated table."""
    table = build_table(load_rubric(input_path), use_name_and_value)
    output_format = output_format.upper()
    if output_format == "CSV":
        write_csv(table, output_path)
    elif output_format in {"EXCEL", "XLSX"}:
        write_excel(table, output_path)
    else:
        raise RubricConversionError(f"Unsupported output format: {output_format}")
    return table
