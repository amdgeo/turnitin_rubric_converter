import csv
import json

import pytest
from openpyxl import load_workbook

from rubric_converter import RubricConversionError, build_table, convert_rubric


@pytest.fixture
def rubric():
    return {
        "RubricScale": [
            {"id": "high", "name": "Excellent"},
            {"id": "low", "name": "Needs work"},
        ],
        "RubricCriterion": [
            {"id": "content", "name": "Content", "description": "Explains the topic", "value": 10}
        ],
        "RubricCriterionScale": [
            {"criterion": "content", "scale_value": "high", "description": "Complete"},
            {"criterion": "content", "scale_value": "low", "description": "Incomplete"},
        ],
        "title": "Example",
    }


def test_build_table_preserves_matrix_and_name_value_option(rubric):
    table = build_table(rubric, use_name_and_value=True)
    assert table.headers == ["Criteria", "Excellent", "Needs work"]
    assert table.rows == [["Content (10)", "Complete", "Incomplete"]]


def test_csv_is_utf8_and_protects_formula_values(tmp_path):
    path = tmp_path / "rubric.rbc"
    path.write_text(json.dumps({
        "RubricScale": [{"id": "s", "name": "Scale"}],
        "RubricCriterion": [{"id": "c", "description": "=unsafe"}],
        "RubricCriterionScale": [{"criterion": "c", "scale_value": "s", "description": "Café"}],
    }), encoding="utf-8")
    output = tmp_path / "out.csv"
    convert_rubric(path, output, "csv")
    with output.open(newline="", encoding="utf-8") as source:
        assert list(csv.reader(source)) == [["Criteria", "Scale"], [["'=unsafe", "Café"]][0]]


def test_excel_contains_matrix_details_and_metadata(tmp_path, rubric):
    source = tmp_path / "rubric.rbc"
    source.write_text(json.dumps(rubric), encoding="utf-8")
    output = tmp_path / "out.xlsx"
    convert_rubric(source, output, "excel")
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Criteria Matrix", "Criteria Details", "Scale Values", "Metadata"]
    assert workbook["Criteria Matrix"]["B2"].value == "Complete"
    assert workbook["Metadata"]["B2"].value == "Example"


def test_excel_handles_list_and_dict_fields_on_criteria_and_scales(tmp_path):
    # Some Turnitin exports include list/dict fields (e.g. related scale IDs)
    # on RubricCriterion/RubricScale items. openpyxl cannot write these
    # directly, so they must be serialized to text instead of raising.
    rubric_with_nested_fields = {
        "RubricScale": [
            {"id": 1, "name": "Excellent", "related_ids": [13098952, 13098953]},
        ],
        "RubricCriterion": [
            {
                "id": 0,
                "description": "Explains the topic",
                "tags": {"category": "content"},
            }
        ],
        "RubricCriterionScale": [
            {"criterion": 0, "scale_value": 1, "description": "Complete"},
        ],
    }
    source = tmp_path / "rubric.rbc"
    source.write_text(json.dumps(rubric_with_nested_fields), encoding="utf-8")
    output = tmp_path / "out.xlsx"
    convert_rubric(source, output, "excel")
    workbook = load_workbook(output)
    details = workbook["Criteria Details"]
    scales = workbook["Scale Values"]
    assert all(
        isinstance(cell.value, (str, int, float, type(None)))
        for row in details.iter_rows()
        for cell in row
    )
    assert all(
        isinstance(cell.value, (str, int, float, type(None)))
        for row in scales.iter_rows()
        for cell in row
    )


def test_unknown_reference_is_rejected(rubric):
    rubric["RubricCriterionScale"][0]["criterion"] = "missing"
    with pytest.raises(RubricConversionError, match="unknown criterion"):
        build_table(rubric)


def test_invalid_json_is_rejected(tmp_path):
    source = tmp_path / "bad.rbc"
    source.write_text("{", encoding="utf-8")
    with pytest.raises(RubricConversionError, match="Could not read rubric"):
        convert_rubric(source, tmp_path / "out.csv", "csv")


def test_integer_ids_are_accepted():
    # Turnitin .rbc exports commonly use numeric IDs rather than strings;
    # these must not be rejected as "missing" IDs.
    rubric_with_integer_ids = {
        "RubricScale": [{"id": 1, "name": "Excellent"}, {"id": 2, "name": "Needs work"}],
        "RubricCriterion": [
            {"id": 0, "name": "Content", "description": "Explains the topic", "value": 10}
        ],
        "RubricCriterionScale": [
            {"criterion": 0, "scale_value": 1, "description": "Complete"},
            {"criterion": 0, "scale_value": 2, "description": "Incomplete"},
        ],
    }
    table = build_table(rubric_with_integer_ids)
    assert table.headers == ["Criteria", "Excellent", "Needs work"]
    assert table.rows == [["Explains the topic", "Complete", "Incomplete"]]
